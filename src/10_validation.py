"""Step 10: manual validation. Draws the random samples, builds the Excel workbook in which
the author records the answers (once), and scores the answers once they are filled in.

  V1 attribution: 9 random issues per tested agent (99 issues). Is the named agent really
     involved in the issue, and in which role? Validates agent_type and the attribution
     signal (D15) on which all three RQs rest.
  V2 topic labels: 5 random issues per fine topic, sampled by 07_rq3_topics.py.
  V3 language: 20 random issues flagged non-English and 20 flagged English (D13).

The sheets are blind: they show what the rater needs (agent named in the dataset, URL,
title, excerpt) but not our classification, which goes to a separate key file. Rows are
shuffled. The answers go into one workbook, docs/validation/VALIDATION_FILLED_BY_AUTHOR.xlsx,
with a tab per sheet and drop-down lists for the answer codes. This script creates the
workbook only if it does not exist yet, so it never overwrites the author's answers.
Scoring: every measure is a share with a Wilson 95% CI, written to
results/tables/validation_summary.csv with a stable measure_id used by the report.
checks/check_08_validation.py re-computes the scores independently. Coding instructions:
docs/validation/README.md.

Reads data/processed/issues_clean.csv and docs/validation/V2_topic_labels_sheet.csv.
Writes docs/validation/V1_attribution_sheet.csv, V3_language_sheet.csv, the workbook (once),
results/tables/validation_V1_key.csv, validation_V3_key.csv and validation_summary.csv.
Run:  python src/10_validation.py
"""
import sys

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from config import MIN_CLOSED_RQ2, PROCESSED, ROOT, SEED, TABLES
from stats_utils import wilson

PER_AGENT = 9  # 11 agents x 9 = 99 issues: about +-10 points at 95% confidence overall
PER_LANGUAGE_GROUP = 20
EXCERPT = 300
VDIR = ROOT / "docs" / "validation"
WORKBOOK = VDIR / "VALIDATION_FILLED_BY_AUTHOR.xlsx"
VDIR.mkdir(parents=True, exist_ok=True)

# answer columns per sheet: (column name, allowed codes or None, help shown when the cell is selected)
ANSWERS = {
    "V1_attribution": [
        (
            "agent_involved (Y/N/U)",
            "Y,N,U",
            "Is this issue really linked to the agent in column B?\nY = yes\nN = no visible link\n"
            "U = cannot tell (page gone, private)",
        ),
        (
            "role (W/T/M/O)",
            "W,T,M,O",
            "Only when Y. How is the agent linked?\nW = the agent wrote it (its bot opened it)\n"
            "T = a person asked the agent to work on it\nM = the agent is only mentioned\n"
            "O = other (explain)",
        ),
        (
            "notes",
            None,
            "Your evidence in a few words, e.g. 'opened by coderabbitai[bot]'. "
            "Always write one for N, U or O.",
        ),
    ],
    "V2_topic_labels": [
        (
            "fits_label (Y/N)",
            "Y,N",
            "Does proposed_label describe what this issue is about?\nY = yes\nN = no\n"
            "Topic 2 has no label: leave empty, describe the issue in notes.",
        ),
        ("notes", None, "Optional; for topic 2, what the issue is about in a few words."),
    ],
    "V3_language": [
        (
            "language (EN/NOT_EN/MIXED)",
            "EN,NOT_EN,MIXED",
            "Is the issue written in English?\nEN = English\nNOT_EN = mainly another language\n"
            "MIXED = large parts in both",
        ),
        ("notes", None, "Optional; for MIXED, which language dominates."),
    ],
}

# Instructions tab: (style, column A, column B, column C); styles: title, head, sub, text, code, example
INSTRUCTIONS = [
    ("title", "Manual validation: how to fill in this workbook", "", ""),
    (
        "text",
        "Why",
        "Our data link every GitHub issue to an AI coding agent with automatic rules, and also flag "
        "its language and its topic. These three checks test by hand whether those rules are right. "
        "Judge only "
        "what you see in the issue; you do not need to know how our data classified it "
        "(that is hidden on purpose).",
        "",
    ),
    (
        "text",
        "Time",
        "V1 about 1.5-2 hours (99 issues), V2 about 30 minutes (25), V3 about 10 minutes (40). "
        "You can stop at any time and continue later.",
        "",
    ),
    ("head", "Steps", "", ""),
    (
        "text",
        "1",
        "Open a tab at the bottom: start with V1_attribution, then V2_topic_labels, then V3_language.",
        "",
    ),
    ("text", "2", "For each row, read the title and the body_excerpt.", ""),
    ("text", "3", "If that is not enough to decide, click 'open issue' to see the issue on GitHub.", ""),
    (
        "text",
        "4",
        "Click a yellow cell and choose your answer from the drop-down arrow. "
        "A small help box explains the codes when you select the cell.",
        "",
    ),
    (
        "text",
        "5",
        "Write a short note with the evidence you saw. Always write one when you answer N, U or O.",
        "",
    ),
    ("text", "6", "Save often (Cmd+S on Mac). Keep the file as .xlsx; do not rename or move it.", ""),
    ("head", "V1_attribution: is the issue really linked to the agent in column B, and how?", "", ""),
    ("sub", "Where to look on GitHub", "", ""),
    (
        "text",
        "a",
        "Who opened the issue: the name at the top of the issue. Is it the agent's own account? "
        "Bot accounts carry a 'bot' tag.",
        "",
    ),
    (
        "text",
        "b",
        "Labels and assignees, on the right-hand side: for example a label 'sweep', or the agent assigned.",
        "",
    ),
    (
        "text",
        "c",
        "The text: '@agent please fix', 'Generated with Claude Code', or the agent only talked about.",
        "",
    ),
    ("sub", "Column agent_involved", "Meaning", "How you recognise it"),
    ("code", "Y", "Yes, the issue is connected to this agent", "Any of the roles W, T, M or O below applies"),
    (
        "code",
        "N",
        "No, there is no visible link to this agent",
        "Nothing in the author, labels, assignees or text " "points to the agent. Explain in notes.",
    ),
    (
        "code",
        "U",
        "Cannot tell",
        "The page is gone (404) or private, or it is genuinely ambiguous. Explain in notes.",
    ),
    ("sub", "Column role (only when Y)", "Meaning", "How you recognise it"),
    (
        "code",
        "W  wrote",
        "The agent itself created the issue",
        "Opened by the agent's own account, usually a bot "
        "(e.g. jules-fleet, coderabbitai[bot], devin-ai-integration[bot]), or the text says it was generated "
        "by or with the agent",
    ),
    (
        "code",
        "T  triggered",
        "A person opened it and asked the agent to work on it",
        "A label such as 'sweep', a title "
        "starting 'Sweep:', '@jules please ...', '/q dev', or the agent is assigned",
    ),
    (
        "code",
        "M  mentioned",
        "A person opened it; the agent is only talked about",
        "A bug report or feature request "
        "about the agent tool itself, or a text or list that names the agent",
    ),
    ("code", "O  other", "Connected in another way", "Explain in notes"),
    ("text", "If several fit", "Choose the strongest: W before T before M.", ""),
    ("sub", "Worked example (this is row 1 of V1; you may copy it)", "", ""),
    (
        "example",
        "What you see",
        "Agent: Jules. Title: '[Fleet Execution] Create batch-generation example'. On GitHub "
        "the issue was opened by the account 'jules-fleet', which carries a bot tag, "
        "and its text says it was "
        "created from a Jules session.",
        "",
    ),
    (
        "example",
        "Your answer",
        "agent_involved = Y;  role = W;  notes = opened by jules-fleet (bot); " "from a Jules session",
        "",
    ),
    ("sub", "More examples (made up, for illustration)", "", ""),
    (
        "example",
        "Y, T",
        "A person opens 'Sweep: fix typo in README' with the label 'sweep'. "
        "Note: person opened; Sweep: prefix and label.",
        "",
    ),
    (
        "example",
        "Y, M",
        "A person reports 'Aider crashes when the repo has submodules'. "
        "Note: bug report about Aider itself.",
        "",
    ),
    (
        "example",
        "N",
        "Agent says Devin, but the issue is a database migration task with no trace of Devin "
        "anywhere. Note: no mention of Devin.",
        "",
    ),
    ("head", "V2_topic_labels: does the proposed label describe the issue?", "", ""),
    ("code", "Y", "The label fits", "Judge what the issue is about as a whole, not single words"),
    ("code", "N", "The label does not fit", "Say in notes what the issue is about instead"),
    (
        "text",
        "Topic 2",
        "Topic 2 has no label. Leave the answer empty and describe each of its issues in notes "
        "in a few words.",
        "",
    ),
    ("head", "V3_language: is the issue written in English?", "", ""),
    (
        "code",
        "EN",
        "English",
        "Title and body are in English (code and English technical words do not count against it)",
    ),
    ("code", "NOT_EN", "Mainly another language", "For example Japanese, Chinese, Korean, Spanish"),
    ("code", "MIXED", "Large parts in both", "Say in notes which language dominates"),
    ("head", "When you are done", "", ""),
    (
        "text",
        "",
        "Save the workbook and tell Claude. The answers are scored with "
        "'python checks/check_08_validation.py'. Full rules: docs/validation/README.md.",
        "",
    ),
]
STYLE = {
    "title": (Font(bold=True, size=16), None),
    "head": (Font(bold=True, size=13, color="FFFFFF"), PatternFill("solid", fgColor="1F4E78")),
    "sub": (Font(bold=True, size=11), PatternFill("solid", fgColor="D9E1F2")),
    "text": (Font(size=11), None),
    "code": (Font(size=11), None),
    "example": (Font(size=11, italic=True), PatternFill("solid", fgColor="F2F2F2")),
}


def excerpt(body):
    return " ".join(str(body).split())[:EXCERPT] if pd.notna(body) else ""


def instructions_tab(ws):
    ws.title = "Instructions"
    for width, letter in zip((22, 70, 70), "ABC"):
        ws.column_dimensions[letter].width = width
    for i, (style, a, b, c) in enumerate(INSTRUCTIONS, start=1):
        font, fill = STYLE[style]
        for j, value in enumerate((a, b, c), start=1):
            cell = ws.cell(row=i, column=j, value=value or None)
            cell.font = Font(bold=True, size=11) if style == "code" and j == 1 else font
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if fill is not None:
                cell.fill = fill
        if style == "head":
            ws.row_dimensions[i].height = 22


def build_workbook(sheets):
    """Instructions tab plus one tab per sheet, with drop-down lists, help boxes and clickable links."""
    wb = Workbook()
    instructions_tab(wb.active)
    answer_fill = PatternFill("solid", fgColor="FFF2CC")
    width = {"item": 6, "topic": 7, "url": 12, "title": 45, "body_excerpt": 70, "notes": 40}
    for name, sheet in sheets.items():
        ws = wb.create_sheet(name)
        answers = {col: (codes, help_text) for col, codes, help_text in ANSWERS[name]}
        columns = list(sheet.columns) + list(answers)
        ws.append(columns)
        for row in sheet.itertuples(index=False):
            ws.append(list(row) + [""] * len(answers))
        n = len(sheet) + 1
        for j, col in enumerate(columns, start=1):
            letter = ws.cell(row=1, column=j).column_letter
            ws.cell(row=1, column=j).font = Font(bold=True)
            ws.column_dimensions[letter].width = width.get(col, 24)
            for r in range(2, n + 1):
                cell = ws.cell(row=r, column=j)
                cell.alignment = Alignment(
                    wrap_text=col in ("title", "body_excerpt", "notes"), vertical="top"
                )
                if col == "url" and cell.value:
                    cell.hyperlink = cell.value
                    cell.value, cell.font = "open issue", Font(color="0563C1", underline="single")
            if col in answers:
                codes, help_text = answers[col]
                for r in range(1, n + 1):
                    ws.cell(row=r, column=j).fill = answer_fill
                dv = DataValidation(
                    type="list" if codes else None,  # None = any value; still shows the help box
                    formula1=f'"{codes}"' if codes else None,
                    allow_blank=True,
                    showInputMessage=True,
                    showErrorMessage=bool(codes),
                )
                dv.promptTitle, dv.prompt = col.split(" (")[0], help_text[:250]
                if codes:
                    dv.errorTitle, dv.error = "Invalid answer", f"Choose one of: {codes}"
                ws.add_data_validation(dv)
                dv.add(f"{letter}2:{letter}{n}")
        ws.freeze_panes = "B2"
    wb.save(WORKBOOK)


d = pd.read_csv(PROCESSED / "issues_clean.csv")
closed_n = d[d.is_closed].agent_type.value_counts()
tested = sorted(closed_n[closed_n >= MIN_CLOSED_RQ2].index)

# V1 attribution: stratified by agent, then shuffled
v1 = pd.concat(d[d.agent_type == a].sample(PER_AGENT, random_state=SEED) for a in tested)
v1 = v1.sample(frac=1, random_state=SEED).reset_index(drop=True)
v1["item"] = np.arange(1, len(v1) + 1)
sheet_v1 = pd.DataFrame(
    {
        "item": v1.item,
        "agent_in_dataset": v1.agent_type,
        "url": v1.html_url,
        "title": v1.title,
        "body_excerpt": v1.body.map(excerpt),
    }
)
key_cols = [
    "item",
    "issue_id",
    "agent_type",
    "author_login",
    "own_bot_author",
    "attribution_signal",
    "s3_precision",
]
v1[key_cols].to_csv(TABLES / "validation_V1_key.csv", index=False)

# V3 language flag: equal samples of flagged non-English and flagged English issues, shuffled
non_en = d[d.lang_group != "english"].sample(PER_LANGUAGE_GROUP, random_state=SEED)
en = d[d.lang_group == "english"].sample(PER_LANGUAGE_GROUP, random_state=SEED)
v3 = pd.concat([non_en, en]).sample(frac=1, random_state=SEED).reset_index(drop=True)
v3["item"] = np.arange(1, len(v3) + 1)
sheet_v3 = pd.DataFrame(
    {"item": v3.item, "url": v3.html_url, "title": v3.title, "body_excerpt": v3.body.map(excerpt)}
)
v3[["item", "issue_id", "agent_type", "lang_group"]].to_csv(TABLES / "validation_V3_key.csv", index=False)

# blank sheets as plain CSV (the published, reproducible samples) and the answer workbook (once)
sheet_v1.to_csv(VDIR / "V1_attribution_sheet.csv", index=False)
sheet_v3.to_csv(VDIR / "V3_language_sheet.csv", index=False)
sheet_v2 = pd.read_csv(VDIR / "V2_topic_labels_sheet.csv").drop(columns=["fits_label (Y/N)", "notes"])
if not WORKBOOK.exists() or "--rebuild-empty-workbook" in sys.argv:
    if WORKBOOK.exists():  # rebuild only a workbook without any answer
        for tab, cols in ANSWERS.items():
            filled = pd.read_excel(WORKBOOK, sheet_name=tab, dtype=str)[[c for c, _, _ in cols]]
            assert filled.isna().all().all(), f"{WORKBOOK.name} already has answers in {tab}; not rebuilt"
    build_workbook({"V1_attribution": sheet_v1, "V2_topic_labels": sheet_v2, "V3_language": sheet_v3})
    print(f"wrote {WORKBOOK.name}")
print(f"V1: {len(v1)} issues ({PER_AGENT} x {len(tested)} agents); V3: {len(v3)} issues")


# -------------------- scoring (skipped while the workbook has no answers)
def answered(tab, column):
    df = pd.read_excel(WORKBOOK, sheet_name=tab, dtype=str).fillna("")
    for col in df.columns:
        df[col] = df[col].str.strip()
    df[column] = df[column].str.upper()
    return df[df[column] != ""]


rows = []


def share(measure_id, description, hits, n):
    pct, lo, hi = wilson(int(hits), int(n)) if n else (np.nan, np.nan, np.nan)
    sheet = measure_id.split("_")[0]
    rows.append(
        {
            "sheet": sheet,
            "measure_id": measure_id,
            "description": description,
            "k": int(hits),
            "n": int(n),
            "pct": pct,
            "ci_lo": lo,
            "ci_hi": hi,
        }
    )


INV, ROLE = "agent_involved (Y/N/U)", "role (W/T/M/O)"
a1 = answered("V1_attribution", INV)
if len(a1):
    a1 = a1.assign(item=a1.item.astype(int)).merge(v1[key_cols], on="item")
    a1[ROLE] = a1[ROLE].str.upper()
    decided = a1[a1[INV].isin(["Y", "N"])]
    share(
        "V1_correct", "attribution correct (Y of Y+N), all agents", (decided[INV] == "Y").sum(), len(decided)
    )
    share("V1_undecidable", "cannot tell (U) of all answered", (a1[INV] == "U").sum(), len(a1))
    for agent, g in decided.groupby("agent_type"):
        share(f"V1_correct_{agent}", f"attribution correct, {agent}", (g[INV] == "Y").sum(), len(g))
    for signal, g in decided.groupby("attribution_signal"):
        share(
            f"V1_correct_signal_{signal}",
            f"attribution correct, signal {signal}",
            (g[INV] == "Y").sum(),
            len(g),
        )
    linked = a1[a1[INV] == "Y"]
    for code, label in [("W", "wrote"), ("T", "triggered"), ("M", "mentioned only"), ("O", "other")]:
        share(
            f"V1_role_{code}",
            f"role {code} ({label}) among linked issues",
            (linked[ROLE] == code).sum(),
            len(linked),
        )
    flagged = a1[a1.own_bot_author & (a1[INV] != "U")]
    share(
        "V1_ownbot_precision",
        "own-bot flag: issue written by the agent (role W)",
        (flagged[ROLE] == "W").sum(),
        len(flagged),
    )
    wrote = linked[linked[ROLE] == "W"]
    share(
        "V1_ownbot_recall",
        "agent-written issues (role W) caught by the own-bot flag",
        wrote.own_bot_author.sum(),
        len(wrote),
    )
    wrong = decided[decided[INV] == "N"]
    share(
        "V1_wrong_other_signal",
        "wrong attributions that come from other (text/branch) signals",
        (wrong.attribution_signal == "other_signal").sum(),
        len(wrong),
    )

a2 = answered("V2_topic_labels", "fits_label (Y/N)")
a2 = a2[a2.proposed_label != "Topic 2"]  # topic 2 had no label to judge
if len(a2):
    share("V2_fit", "label fits, all labelled fine topics", (a2["fits_label (Y/N)"] == "Y").sum(), len(a2))
    for topic, g in a2.groupby("topic"):
        share(
            f"V2_fit_T{topic}",
            f"label fits, fine topic {topic}",
            (g["fits_label (Y/N)"] == "Y").sum(),
            len(g),
        )

LANG = "language (EN/NOT_EN/MIXED)"
a3 = answered("V3_language", LANG)
if len(a3):
    a3 = a3.assign(item=a3.item.astype(int)).merge(v3[["item", "lang_group"]], on="item")
    flagged_non, flagged_en = a3[a3.lang_group != "english"], a3[a3.lang_group == "english"]
    share(
        "V3_nonEN_correct",
        "flagged non-English that is not English",
        (flagged_non[LANG] != "EN").sum(),
        len(flagged_non),
    )
    share(
        "V3_EN_correct", "flagged English that is English", (flagged_en[LANG] == "EN").sum(), len(flagged_en)
    )

summary = pd.DataFrame(
    rows, columns=["sheet", "measure_id", "description", "k", "n", "pct", "ci_lo", "ci_hi"]
)
summary.to_csv(TABLES / "validation_summary.csv", index=False, float_format="%.6g")
print(f"validation answers: V1 {len(a1)}, V2 {len(a2)}, V3 {len(a3)}; {len(summary)} measures scored")
